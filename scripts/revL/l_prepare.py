from pathlib import Path
from collections import Counter
import sys, json, re
import openpyxl
from pypdf import PdfReader
import pypdfium2 as pdfium

sys.stdout.reconfigure(encoding='utf-8')
R=Path(__file__).resolve().parents[2]
SRC=R/'work/intel_datasheets'
O=R/'work/revL'
O.mkdir(parents=True,exist_ok=True)

def norm(v):
    if v is None or isinstance(v,(str,int,float,bool)):return v
    return str(v)

all_sheets={}
for f in SRC.glob('*.xlsx'):
    w=openpyxl.load_workbook(f,read_only=True,data_only=True)
    all_sheets[f.name]={s.title:[[norm(v) for v in row] for row in s.iter_rows(values_only=True)] for s in w}
(O/'intel_tables.json').write_text(json.dumps(all_sheets,ensure_ascii=False,indent=2),encoding='utf-8')
rows=all_sheets['825698-001_Ballout.xlsx']['Pinlist']
pins=[]
for i,row in enumerate(rows[1:],2):
    if row[0] is None:continue
    if not re.fullmatch(r'[A-Z]{1,2}\d{1,2}',str(row[0])):continue
    pins.append(dict(ball=row[0],signal=row[1],io=row[2],interface=row[3],x_um=row[4],y_um=row[5],comment=row[6],source_row=i))
assert len(pins)==2833,len(pins)
assert len({q['ball'] for q in pins})==2833
assert len({(q['x_um'],q['y_um']) for q in pins})==2833
out={'source':'Intel 829568-004, embedded 825698-001_Ballout.xlsx, Pinlist',
     'coordinate_unit':'um; exact manufacturer coordinates retained; orientation to be verified',
     'count':len(pins),'pins':pins}
(O/'INTEL_268V_PINOUT.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print('ball count',len(pins),'interfaces',dict(Counter(q['interface'] for q in pins)))
print('bounds um',min(q['x_um'] for q in pins),max(q['x_um'] for q in pins),min(q['y_um'] for q in pins),max(q['y_um'] for q in pins))
for term in ['PCI','SVID','Power','Clock','SPI','RTC','Reset']:
    pp=[q for q in pins if term.lower() in (q['interface'] or '').lower() or term.lower() in (q['signal'] or '').lower()]
    print(term,len(pp),[(q['ball'],q['signal']) for q in pp[:55]])

f=SRC/'Intel_Core_Ultra_200V_829568-004.pdf'
r=PdfReader(f)
(O/'intel_829568.txt').write_text('\n\n'.join('PAGE '+str(i+1)+'\n'+p.extract_text() for i,p in enumerate(r.pages)),encoding='utf-8')
doc=pdfium.PdfDocument(str(f))
for page in [19,25,36,37,120,121,123,140,141,142,143,144,158,159,160,161,162,163,164]:
    doc[page-1].render(scale=1.5).to_pil().save(O/f'intel_page_{page}.png')
print('extracted source tables, text, and key page images')
